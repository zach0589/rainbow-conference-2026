#!/usr/bin/env python3
"""Build Rainbow Conference 2026 Media Mix PPTX from scratch."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import os

# ── Constants ──
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
NAVY = RGBColor(0x1B, 0x30, 0x57)
RED = RGBColor(0xC8, 0x10, 0x2E)
GRAY = RGBColor(0x58, 0x59, 0x5B)
LIGHT_GRAY_BG = RGBColor(0xF2, 0xF2, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_RED = RGBColor(0x9B, 0x1B, 0x30)
SOURCE_GRAY = RGBColor(0xAA, 0xAA, 0xAA)
FONT = 'Segoe UI'

ASSET_DIR = '/private/tmp/rainbow-conference-2026/assets'
RAINBOW_LOGO = os.path.join(ASSET_DIR, 'rainbow-logo.png')
SOS_LOGO = os.path.join(ASSET_DIR, 'sos-logo.png')
SOS_LOGO_WHITE = os.path.join(ASSET_DIR, 'sos-logo-white.png')

# Padding
PAD_L = Inches(0.6)
PAD_R = Inches(0.6)
PAD_T = Inches(0.5)
CONTENT_W = SLIDE_W - PAD_L - PAD_R


def set_font(run, size=13, color=GRAY, bold=False, name=FONT):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = name


def add_textbox(slide, left, top, width, height, text, size=13, color=GRAY,
                bold=False, alignment=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    # Set vertical anchor
    txBox.text_frame._txBody.bodyPr.set('anchor', {
        MSO_ANCHOR.TOP: 't', MSO_ANCHOR.MIDDLE: 'ctr', MSO_ANCHOR.BOTTOM: 'b'
    }.get(anchor, 't'))
    p = tf.paragraphs[0]
    p.alignment = alignment
    run = p.add_run()
    run.text = text
    set_font(run, size, color, bold)
    return txBox


def add_multiline_textbox(slide, left, top, width, height, lines, alignment=PP_ALIGN.LEFT, anchor='t'):
    """lines: list of (text, size, color, bold) tuples or list of list of (text, size, color, bold) for mixed runs."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    txBox.text_frame._txBody.bodyPr.set('anchor', anchor)
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = alignment
        p.space_before = Pt(0)
        p.space_after = Pt(2)
        if isinstance(line, list):
            # Multiple runs in one paragraph
            for text, size, color, bold in line:
                run = p.add_run()
                run.text = text
                set_font(run, size, color, bold)
        else:
            text, size, color, bold = line
            run = p.add_run()
            run.text = text
            set_font(run, size, color, bold)
    return txBox


def add_logos(slide, white_bg=True):
    """Add logos to bottom corners."""
    rainbow_h = Inches(0.45)
    sos_h = Inches(0.3)
    # Rainbow logo: bottom-left
    slide.shapes.add_picture(RAINBOW_LOGO,
                             Inches(0.4),
                             SLIDE_H - Inches(0.15) - rainbow_h,
                             height=rainbow_h)
    # SOS logo: bottom-right
    sos_path = SOS_LOGO if white_bg else SOS_LOGO_WHITE
    slide.shapes.add_picture(sos_path,
                             SLIDE_W - Inches(0.15) - Inches(0.9),
                             SLIDE_H - Inches(0.2) - sos_h,
                             height=sos_h)


def add_red_underline(slide, left, top, width):
    """Add a 3pt red line."""
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(3))
    line.fill.solid()
    line.fill.fore_color.rgb = RED
    line.line.fill.background()
    return line


def add_title_with_underline(slide, text, top=PAD_T, left=PAD_L, width=None):
    """Standard slide title: 20pt navy bold with red underline."""
    if width is None:
        width = CONTENT_W
    tb = add_textbox(slide, left, top, width, Inches(0.45), text, 20, NAVY, True)
    add_red_underline(slide, left, top + Inches(0.42), width)
    return top + Inches(0.55)


def add_bullet_list(slide, items, left, top, width, height, size=13, color=GRAY):
    """Add a bulleted text box. Items: list of str or list of (runs) for mixed formatting."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(2)
        p.space_after = Pt(4)
        # Add bullet
        pPr = p._pPr if p._pPr is not None else p._p.get_or_add_pPr()
        buNone = pPr.find(qn('a:buNone'))
        if buNone is not None:
            pPr.remove(buNone)
        buChar = pPr.makeelement(qn('a:buChar'), {'char': '\u2022'})
        pPr.append(buChar)
        buClr = pPr.makeelement(qn('a:buClr'), {})
        srgb = buClr.makeelement(qn('a:srgbClr'), {'val': 'C8102E'})
        buClr.append(srgb)
        pPr.append(buClr)
        buSzPct = pPr.makeelement(qn('a:buSzPct'), {'val': '100000'})
        pPr.append(buSzPct)
        buFont = pPr.makeelement(qn('a:buFont'), {'typeface': FONT})
        pPr.append(buFont)
        p.level = 0
        pPr.set('marL', str(Emu(Inches(0.25))))
        pPr.set('indent', str(Emu(-Inches(0.2))))

        if isinstance(item, list):
            for text, sz, clr, bld in item:
                run = p.add_run()
                run.text = text
                set_font(run, sz, clr, bld)
        else:
            run = p.add_run()
            run.text = item
            set_font(run, size, color, False)
    return txBox


def add_callout(slide, text, left, top, width, height=Inches(0.5),
                border_color=NAVY, bg_color=None, text_parts=None):
    """Callout box with left border."""
    if bg_color is None:
        bg_color = RGBColor(0xED, 0xEE, 0xF0) if border_color == NAVY else RGBColor(0xFD, 0xEE, 0xF0)
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = bg_color
    box.line.fill.background()
    # Left border
    border = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Pt(3), height)
    border.fill.solid()
    border.fill.fore_color.rgb = border_color
    border.line.fill.background()
    # Text
    txBox = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.05),
                                      width - Inches(0.2), height - Inches(0.1))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    if text_parts:
        for t, sz, clr, bld in text_parts:
            run = p.add_run()
            run.text = t
            set_font(run, sz, clr, bld)
    else:
        run = p.add_run()
        run.text = text
        set_font(run, 11, GRAY, False)
    return box


def add_stat_card(slide, left, top, width, height, number, label,
                  border_color=RED, num_size=28):
    """Stat card: rounded rect with colored top border."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = LIGHT_GRAY_BG
    card.line.fill.background()
    card.adjustments[0] = 0.05
    # Top border
    border = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(4))
    border.fill.solid()
    border.fill.fore_color.rgb = border_color
    border.line.fill.background()
    # Number
    add_textbox(slide, left, top + Inches(0.15), width, Inches(0.45),
                number, num_size, NAVY, True, PP_ALIGN.CENTER)
    # Label
    add_textbox(slide, left + Inches(0.05), top + Inches(0.55), width - Inches(0.1), Inches(0.5),
                label, 10, GRAY, False, PP_ALIGN.CENTER)
    return card


def add_source(slide, text, top=None):
    """Source line at bottom."""
    if top is None:
        top = SLIDE_H - Inches(0.65)
    add_textbox(slide, Inches(0), top, SLIDE_W, Inches(0.25),
                text, 8, SOURCE_GRAY, False, PP_ALIGN.CENTER)


def add_bg_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def set_slide_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_freeform_trapezoid(slide, points, fill_color, opacity_pct=92):
    """Add a freeform polygon from list of (x_emu, y_emu) points (pass Inches values)."""
    builder = slide.shapes.build_freeform(points[0][0], points[0][1])
    # add_line_segments takes a sequence of (x, y) vertices
    builder.add_line_segments(points[1:] + [points[0]])
    shape = builder.convert_to_shape()
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    # Set opacity via XML
    alpha_val = int((opacity_pct / 100) * 100000)
    # Access the spPr element directly
    spPr = shape._element.spPr
    solidFill = spPr.find(qn('a:solidFill'))
    if solidFill is not None:
        srgb = solidFill.find(qn('a:srgbClr'))
        if srgb is not None:
            alpha_elem = srgb.makeelement(qn('a:alpha'), {'val': str(alpha_val)})
            srgb.append(alpha_elem)
    return shape


# ══════════════════════════════════════════════════════════════
# BUILD PRESENTATION
# ══════════════════════════════════════════════════════════════

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
blank_layout = prs.slide_layouts[6]  # Blank layout

# ── SLIDE 1: Title ──
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, NAVY)
tb = add_multiline_textbox(slide, Inches(1), Inches(1.5), Inches(11.333), Inches(2.5), [
    ("The 2026 Restoration", 36, WHITE, True),
    ("Marketing Reality", 36, WHITE, True),
], PP_ALIGN.CENTER, 'ctr')
add_textbox(slide, Inches(1), Inches(4.0), Inches(11.333), Inches(0.5),
            "Winning with the Right Media Mix", 18, RED, False, PP_ALIGN.CENTER)
add_textbox(slide, Inches(1), Inches(6.2), Inches(11.333), Inches(0.35),
            "Rainbow Restoration National Conference 2026", 11,
            RGBColor(0x66, 0x66, 0x66), False, PP_ALIGN.CENTER)
add_logos(slide, white_bg=False)

# ── SLIDE 2: Split Statement ──
slide = prs.slides.add_slide(blank_layout)
add_bg_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(3.75), WHITE)
add_bg_rect(slide, Inches(0), Inches(3.75), SLIDE_W, Inches(3.75), NAVY)
line_w = Inches(13.333 * 0.8)
line_l = (SLIDE_W - line_w) // 2
add_red_underline(slide, line_l, Inches(3.74), line_w)
add_textbox(slide, Inches(1), Inches(1.0), Inches(11.333), Inches(2.5),
            "How your customers find you has changed", 32, NAVY, True, PP_ALIGN.CENTER,
            MSO_ANCHOR.BOTTOM)
add_textbox(slide, Inches(1), Inches(4.0), Inches(11.333), Inches(2.5),
            "Has your marketing kept up?", 20, RGBColor(0x99, 0x99, 0x99), False,
            PP_ALIGN.CENTER, MSO_ANCHOR.TOP)
add_logos(slide, white_bg=False)

# ── SLIDE 3: Media Mix ──
slide = prs.slides.add_slide(blank_layout)
y = add_title_with_underline(slide, "Your Media Mix Is Your System")
add_bullet_list(slide, [
    "It's every marketing channel you invest in \u2014 and how they connect",
    [("Most operators have a list of channels. The best operators have a ", 13, GRAY, False),
     ("system.", 13, GRAY, True)],
], PAD_L, y, CONTENT_W, Inches(0.7))
y += Inches(0.85)
box_w = (CONTENT_W - Inches(0.3)) / 3
box_h = Inches(1.3)
box_gap = Inches(0.15)
box_configs = [
    ("CAPTURE", "Demand that\nalready exists", NAVY),
    ("CREATE", "Demand before the\nemergency happens", RED),
    ("CONVERT", "Attention into\nbooked jobs", GRAY),
]
for i, (title, subtitle, color) in enumerate(box_configs):
    bx = PAD_L + i * (box_w + box_gap)
    bg_clr = RGBColor(0xED, 0xEE, 0xF0) if color == NAVY else (
        RGBColor(0xFD, 0xEE, 0xF0) if color == RED else RGBColor(0xEE, 0xEE, 0xEF))
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, y, box_w, box_h)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_clr
    card.line.fill.background()
    card.adjustments[0] = 0.04
    border = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, bx, y, box_w, Pt(4))
    border.fill.solid()
    border.fill.fore_color.rgb = color
    border.line.fill.background()
    add_textbox(slide, bx, y + Inches(0.2), box_w, Inches(0.4),
                title, 22, color, True, PP_ALIGN.CENTER)
    add_textbox(slide, bx + Inches(0.1), y + Inches(0.65), box_w - Inches(0.2), Inches(0.5),
                subtitle, 10, GRAY, False, PP_ALIGN.CENTER)
y += box_h + Inches(0.15)
add_callout(slide, "Most restoration companies are good at one of these. The ones growing are covering all three.",
            PAD_L, y, CONTENT_W, Inches(0.45))
add_logos(slide)

# ── SLIDE 4: Funnel ──
slide = prs.slides.add_slide(blank_layout)
y = add_title_with_underline(slide, "The Marketing Funnel")
funnel_left = PAD_L
funnel_w_base = Inches(4.0)
tier_h = Inches(0.9)
gap = Inches(0.07)
funnel_data = [
    ("AWARENESS", RED, 92, 1.0, 0.85),
    ("CONSIDERATION", DARK_RED, 88, 0.85, 0.65),
    ("DECISION", NAVY, 92, 0.65, 0.42),
    ("CONVERSION", GRAY, 88, 0.42, 0.25),
]
fy = y + Inches(0.1)
for i, (label, color, opacity, top_pct, bot_pct) in enumerate(funnel_data):
    cx = funnel_left + funnel_w_base / 2
    tw = funnel_w_base * top_pct
    bw = funnel_w_base * bot_pct
    tl = cx - tw / 2
    tr = cx + tw / 2
    bl = cx - bw / 2
    br = cx + bw / 2
    top_y = fy + i * (tier_h + gap)
    bot_y = top_y + tier_h
    shape = add_freeform_trapezoid(slide, [
        (tl, top_y), (tr, top_y), (br, bot_y), (bl, bot_y)
    ], color, opacity)
    mid_y = top_y + tier_h / 2 - Inches(0.15)
    mid_w = (tw + bw) / 2
    mid_l = cx - mid_w / 2
    add_textbox(slide, mid_l, mid_y, mid_w, Inches(0.3),
                label, 13, WHITE, True, PP_ALIGN.CENTER)

label_left = PAD_L + funnel_w_base + Inches(0.4)
label_w = CONTENT_W - funnel_w_base - Inches(0.4)
funnel_labels = [
    ("Awareness", RED, "Social media, sponsorships, YouTube, vehicle wraps"),
    ("Consideration", DARK_RED, "Retargeting, blog content, SEO, email nurture"),
    ("Decision", NAVY, "LSA, Google Ads, SEO (high-intent), GBP"),
    ("Conversion", GRAY, "Website, reviews, speed to lead, follow-up"),
]
for i, (name, color, desc) in enumerate(funnel_labels):
    row_y = fy + i * (tier_h + gap) + tier_h / 2 - Inches(0.25)
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, label_left, row_y + Inches(0.08),
                                  Inches(0.12), Inches(0.12))
    dot.fill.solid()
    dot.fill.fore_color.rgb = color
    dot.line.fill.background()
    add_multiline_textbox(slide, label_left + Inches(0.2), row_y - Inches(0.02),
                          label_w - Inches(0.2), Inches(0.5), [
        (name, 13, color, True),
        (desc, 10, GRAY, False),
    ])
callout_y = fy + 4 * (tier_h + gap) + Inches(0.15)
add_callout(slide,
            None, PAD_L, callout_y, CONTENT_W, Inches(0.5),
            RED, RGBColor(0xFD, 0xEE, 0xF0),
            text_parts=[
                ("Restoration is emergency-driven. ", 11, RED, True),
                ("The consideration phase is far shorter than most industries. Trust has to be built BEFORE the emergency.", 11, GRAY, False),
            ])
add_logos(slide)

# ── SLIDE 5: Landscape ──
slide = prs.slides.add_slide(blank_layout)
y = add_title_with_underline(slide, "The Landscape Is Shifting")
card_w = (CONTENT_W - Inches(0.3)) / 3
card_h = Inches(1.1)
card_gap = Inches(0.15)
stats5 = [
    ("29%", "of homeowners already have a\ncompany name in mind\nwhen they search", RED),
    ("+27.8%", "Google Ads spend growth\nacross our restoration clients", NAVY),
    ("+63.9%", "LSA spend growth\nacross our restoration clients", RED),
]
for i, (num, label, bc) in enumerate(stats5):
    cx = PAD_L + i * (card_w + card_gap)
    add_stat_card(slide, cx, y, card_w, card_h, num, label, bc)
y += card_h + Inches(0.15)
add_bullet_list(slide, [
    [("Businesses in LSA results get ", 11, GRAY, False),
     ("25-30% more calls", 11, GRAY, True),
     (" than those relying on organic alone", 11, GRAY, False)],
    [("31% of consumers now require 4.5+ stars", 11, GRAY, True),
     (" to consider a business \u2014 up from 17% last year", 11, GRAY, False)],
], PAD_L, y, CONTENT_W, Inches(0.65))
y += Inches(0.7)
add_callout(slide, "Your competitors are spending more every quarter. But the operators winning aren't just spending \u2014 they're building recognition before the emergency.",
            PAD_L, y, CONTENT_W, Inches(0.5))
add_source(slide, "Sources: SOS x C&R Magazine Report 2026; AdMall/SalesFuel 2025; BrightLocal LCRS 2026; Boomcycle 2026")
add_logos(slide)

# ── SLIDE 6: Customers ──
slide = prs.slides.add_slide(blank_layout)
y = add_title_with_underline(slide, "What Your Customers Respond To")
stats6 = [
    ("97%", "of restoration operators say referrals\nare their #1 source of new business", RED),
    ("58%", "of homeowners were influenced\nby a blog post in the last 30 days", NAVY),
    ("55,352", "restoration establishments\ncompeting for attention nationwide", RED),
]
for i, (num, label, bc) in enumerate(stats6):
    cx = PAD_L + i * (card_w + card_gap)
    add_stat_card(slide, cx, y, card_w, card_h, num, label, bc)
y += card_h + Inches(0.15)
add_bullet_list(slide, [
    "Homeowners are 12% more likely to check a company's blog to evaluate expertise",
    [("41% called or visited the advertiser after seeing an ad \u2014 ", 11, GRAY, False),
     ("136% higher", 11, GRAY, True),
     (" than the general population", 11, GRAY, False)],
    [("Only 3.5% of restoration customers say they don't read online reviews \u2014 ", 11, GRAY, False),
     ("nearly everyone checks", 11, GRAY, True)],
], PAD_L, y, CONTENT_W, Inches(0.75))
add_source(slide, "Sources: AdMall/SalesFuel 2025; Cleanfax Restoration Benchmarking Survey 2024")
add_logos(slide)

# ── SLIDE 7: AI ──
slide = prs.slides.add_slide(blank_layout)
y = add_title_with_underline(slide, "AI Is Changing How Customers Find You")
stats7 = [
    ("14.2%", "AI Conversion Rate", RED),
    ("2.8%", "Google Conversion Rate", NAVY),
    ("5x", "AI Converts Higher", GRAY),
]
for i, (num, label, bc) in enumerate(stats7):
    cx = PAD_L + i * (card_w + card_gap)
    add_stat_card(slide, cx, y, card_w, card_h, num, label, bc)
y += card_h + Inches(0.2)
col_w = (CONTENT_W - Inches(0.3)) / 2
add_textbox(slide, PAD_L, y, col_w, Inches(0.25),
            "THE REALITY", 10, NAVY, True)
add_bullet_list(slide, [
    "Google still handles 90% of search",
    [("AI converts at ", 11, GRAY, False), ("5x the rate", 11, GRAY, True)],
    "41% of restoration customers use AI to research",
], PAD_L, y + Inches(0.25), col_w, Inches(0.8), 11)
right_x = PAD_L + col_w + Inches(0.3)
add_textbox(slide, right_x, y, col_w, Inches(0.25),
            "THE GOOD NEWS", 10, NAVY, True)
add_bullet_list(slide, [
    [("Google and AI reward the ", 11, GRAY, False), ("same thing", 11, GRAY, True)],
    "Authoritative content, strong reviews, managed GBP",
    "A brand worth citing \u2014 AI recommends what it \"knows\"",
], right_x, y + Inches(0.25), col_w, Inches(0.8), 11)
y += Inches(1.2)
add_callout(slide, None, PAD_L, y, CONTENT_W, Inches(0.45),
            RED, RGBColor(0xFD, 0xEE, 0xF0),
            text_parts=[
                ("You don't need two strategies. ", 11, RED, True),
                ("Google and AI are two systems rewarding the same thing.", 11, GRAY, False),
            ])
add_source(slide, "Source: Superprompt, 12M Visit Study, 2025")
add_logos(slide)

# ── SLIDE 8: What This Means ──
slide = prs.slides.add_slide(blank_layout)
y = add_title_with_underline(slide, "What This Means for Your Funnel")
items8 = [
    [("The bottom of the funnel still matters most. ", 13, RED, True),
     ("LSA and PPC are where leads come from. But costs are up 20.9% \u2014 the bottom alone isn't enough anymore.", 13, GRAY, False)],
    [("The top of the funnel makes the bottom work harder. ", 13, RED, True),
     ("Branded searches convert at 6x the rate of generic terms. Awareness doesn't replace demand capture \u2014 it supercharges it.", 13, GRAY, False)],
    [("AI is the new front door. ", 13, RED, True),
     ("41% of your potential customers are using AI to research you. Your content, reviews, and GBP feed what AI recommends.", 13, GRAY, False)],
]
add_bullet_list(slide, items8, PAD_L, y, CONTENT_W, Inches(2.5))
y += Inches(2.7)
add_callout(slide, "Keep investing in the bottom. But the operators pulling ahead are the ones adding layers above it.",
            PAD_L, y, CONTENT_W, Inches(0.45), NAVY)
add_source(slide, "Sources: SOS x C&R Magazine Report 2026; AdMall/SalesFuel 2025")
add_logos(slide)

# ── SLIDE 9: Workshop Divider ──
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, RED)
add_textbox(slide, Inches(1), Inches(2.0), Inches(11.333), Inches(1.0),
            "WORKSHOP", 44, WHITE, True, PP_ALIGN.CENTER)
add_textbox(slide, Inches(1), Inches(3.0), Inches(11.333), Inches(0.5),
            "Build Your Marketing Funnel", 18,
            RGBColor(0xFF, 0xCC, 0xCC), False, PP_ALIGN.CENTER)
qr_size = Inches(1.0)
qr_left = Inches(5.0)
qr_top = Inches(4.0)
qr_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, qr_left, qr_top, qr_size, qr_size)
qr_box.fill.solid()
qr_box.fill.fore_color.rgb = WHITE
qr_box.line.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
qr_box.adjustments[0] = 0.08
add_textbox(slide, qr_left - Inches(0.5), qr_top + qr_size + Inches(0.1),
            qr_size + Inches(1.0), Inches(0.35),
            "QR CODE", 8, RGBColor(0xCC, 0xCC, 0xCC), False, PP_ALIGN.CENTER)
add_multiline_textbox(slide, qr_left + qr_size + Inches(0.2), qr_top + Inches(0.05),
                      Inches(3.5), Inches(0.9), [
    ("Scan to open the", 12, RGBColor(0xFF, 0xDD, 0xDD), False),
    ("Workshop Worksheet", 12, WHITE, True),
    ("or use the printed copy at your seat", 10, RGBColor(0xFF, 0xAA, 0xAA), False),
])
add_logos(slide, white_bg=False)

# ── SLIDE 10: Funnel Reference ──
slide = prs.slides.add_slide(blank_layout)
add_textbox(slide, Inches(1), Inches(0.5), Inches(11.333), Inches(0.5),
            "Where Does Your Marketing Fit?", 22, NAVY, True, PP_ALIGN.CENTER)
add_textbox(slide, Inches(1), Inches(1.0), Inches(11.333), Inches(0.35),
            "Map each activity to one of these four stages.", 12, GRAY, False, PP_ALIGN.CENTER)
funnel_w = Inches(5.0)
funnel_cx = SLIDE_W / 2
funnel_top = Inches(1.6)
tier_h10 = Inches(1.15)
gap10 = Inches(0.08)
funnel_data10 = [
    ("AWARENESS", RED, 92, 1.0, 0.82),
    ("CONSIDERATION", DARK_RED, 88, 0.82, 0.60),
    ("DECISION", NAVY, 92, 0.60, 0.38),
    ("CONVERSION", GRAY, 88, 0.38, 0.20),
]
for i, (label, color, opacity, top_pct, bot_pct) in enumerate(funnel_data10):
    tw = funnel_w * top_pct
    bw = funnel_w * bot_pct
    tl = funnel_cx - tw / 2
    tr = funnel_cx + tw / 2
    bl = funnel_cx - bw / 2
    br = funnel_cx + bw / 2
    ty = funnel_top + i * (tier_h10 + gap10)
    by = ty + tier_h10
    shape = add_freeform_trapezoid(slide, [
        (tl, ty), (tr, ty), (br, by), (bl, by)
    ], color, opacity)
    mid_y = ty + tier_h10 / 2 - Inches(0.12)
    mid_w = (tw + bw) / 2
    add_textbox(slide, funnel_cx - mid_w / 2, mid_y, mid_w, Inches(0.3),
                label, 16, WHITE, True, PP_ALIGN.CENTER)
add_logos(slide)

# ── SLIDE 11: Step 1 ──
slide = prs.slides.add_slide(blank_layout)
border_rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                      Inches(0.4), Inches(0.35),
                                      SLIDE_W - Inches(0.8), SLIDE_H - Inches(1.0))
border_rect.fill.background()
border_rect.line.color.rgb = RED
border_rect.line.width = Pt(3)
border_rect.adjustments[0] = 0.02
add_textbox(slide, Inches(0.9), Inches(0.8), Inches(11), Inches(0.45),
            "Step 1: Map What You're Doing", 20, RED, True)
items11 = [
    "Write down every marketing activity you're currently doing",
    [("For each one, answer: ", 13, GRAY, False), ("WHY", 13, GRAY, True),
     (" are you doing this?", 13, GRAY, False)],
    [("Then answer: ", 13, GRAY, False), ("WHERE", 13, GRAY, True),
     (" does it fit in the funnel?", 13, GRAY, False)],
    "Don't overthink it. Just get it all on paper.",
]
add_bullet_list(slide, items11, Inches(0.9), Inches(1.5), Inches(11), Inches(3.5))
add_logos(slide)

# ── SLIDE 12: Step 2 ──
slide = prs.slides.add_slide(blank_layout)
border_rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                      Inches(0.4), Inches(0.35),
                                      SLIDE_W - Inches(0.8), SLIDE_H - Inches(1.0))
border_rect.fill.background()
border_rect.line.color.rgb = RED
border_rect.line.width = Pt(3)
border_rect.adjustments[0] = 0.02
add_textbox(slide, Inches(0.9), Inches(0.8), Inches(11), Inches(0.45),
            "Step 2: Find the Gaps", 20, RED, True)
items12 = [
    "Which funnel stages have the most activity? Which have none?",
    "Heavy on Decision (LSA, PPC) but light on everything else?",
    "No Awareness activity at all?",
    'Nothing in Conversion beyond "we answer the phone"?',
    [("Circle the biggest gap \u2014 that's your #1 priority", 13, RED, True)],
]
add_bullet_list(slide, items12, Inches(0.9), Inches(1.5), Inches(11), Inches(3.5))
add_logos(slide)

# ── SLIDE 13: Step 3 ──
slide = prs.slides.add_slide(blank_layout)
border_rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                      Inches(0.4), Inches(0.35),
                                      SLIDE_W - Inches(0.8), SLIDE_H - Inches(1.0))
border_rect.fill.background()
border_rect.line.color.rgb = RED
border_rect.line.width = Pt(3)
border_rect.adjustments[0] = 0.02
add_textbox(slide, Inches(0.9), Inches(0.8), Inches(11), Inches(0.45),
            "Step 3: Pick Your Next Move", 20, RED, True)
items13 = [
    [("Based on your gaps, choose ", 13, GRAY, False), ("1-3 things", 13, GRAY, True),
     (" you want to add to your marketing", 13, GRAY, False)],
    [("For each one, write down: ", 13, GRAY, False), ("What is it?", 13, GRAY, True)],
    "Where does it fit in the funnel?",
    "Why this one?",
    [("Don't worry about the \"how\" yet \u2014 ", 13, GRAY, False),
     ("that's what we're here for", 13, RED, True)],
]
add_bullet_list(slide, items13, Inches(0.9), Inches(1.5), Inches(11), Inches(3.5))
add_logos(slide)

# ── SLIDE 14: Close ──
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, NAVY)
logo_y = Inches(1.0)
slide.shapes.add_picture(RAINBOW_LOGO,
                         Inches(4.5), logo_y,
                         height=Inches(1.0))
slide.shapes.add_picture(SOS_LOGO_WHITE,
                         Inches(7.5), logo_y + Inches(0.15),
                         height=Inches(0.65))
add_multiline_textbox(slide, Inches(1), Inches(2.5), Inches(11.333), Inches(2.5), [
    ("Build the System.", 30, WHITE, True),
    ("Cover the Funnel.", 30, WHITE, True),
    ("Win the Market.", 30, WHITE, True),
], PP_ALIGN.CENTER, 'ctr')
add_textbox(slide, Inches(1), Inches(5.2), Inches(11.333), Inches(0.5),
            "Authoritative content. Strong reviews. A brand worth citing.",
            14, RGBColor(0x88, 0x99, 0xAA), False, PP_ALIGN.CENTER)

# ── SLIDE 15: QR Codes ──
slide = prs.slides.add_slide(blank_layout)
set_slide_bg(slide, NAVY)
logo_y = Inches(0.6)
slide.shapes.add_picture(RAINBOW_LOGO,
                         Inches(4.5), logo_y,
                         height=Inches(1.0))
slide.shapes.add_picture(SOS_LOGO_WHITE,
                         Inches(7.5), logo_y + Inches(0.15),
                         height=Inches(0.65))
add_textbox(slide, Inches(1), Inches(2.0), Inches(11.333), Inches(0.5),
            "Scan Before You Leave", 24, WHITE, False, PP_ALIGN.CENTER)
qr_size = Inches(1.5)
qr_y = Inches(3.0)
qr_gap = Inches(1.5)
total_qr_w = 3 * qr_size + 2 * qr_gap
qr_start = (SLIDE_W - total_qr_w) / 2
qr_labels = ["Download the Report", "Workshop Worksheet", "Talk to an SOS\nMarketing Consultant"]
for i in range(3):
    qx = qr_start + i * (qr_size + qr_gap)
    qr = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, qx, qr_y, qr_size, qr_size)
    qr.fill.solid()
    qr.fill.fore_color.rgb = WHITE
    qr.line.fill.background()
    qr.adjustments[0] = 0.06
    add_textbox(slide, qx, qr_y + Inches(0.5), qr_size, Inches(0.3),
                "QR CODE", 9, RGBColor(0xCC, 0xCC, 0xCC), False, PP_ALIGN.CENTER)
    add_textbox(slide, qx - Inches(0.25), qr_y + qr_size + Inches(0.15),
                qr_size + Inches(0.5), Inches(0.5),
                qr_labels[i], 11, WHITE, True, PP_ALIGN.CENTER)

# ── Verify and save ──
assert len(prs.slides) == 15, f"Expected 15 slides, got {len(prs.slides)}"

output_path = '/private/tmp/rainbow-conference-2026/Rainbow_Conference_2026_Media_Mix.pptx'
prs.save(output_path)
print(f"Saved {output_path} with {len(prs.slides)} slides")
